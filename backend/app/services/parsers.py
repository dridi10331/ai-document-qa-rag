from dataclasses import dataclass
from pathlib import Path
import re

from PyPDF2 import PdfReader
from docx import Document as DocxDocument

from app.services.ocr import ocr_pdf_page

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class ParsedPage:
    text: str
    page_number: int | None
    metadata: dict


@dataclass
class ParsedDocument:
    pages: list[ParsedPage]
    metadata: dict


def parse_pdf(file_path: Path, enable_ocr: bool) -> ParsedDocument:
    reader = PdfReader(str(file_path))
    pages: list[ParsedPage] = []
    for index, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if enable_ocr and not text.strip():
            text = ocr_pdf_page(file_path, index)
        pages.append(
            ParsedPage(
                text=text,
                page_number=index + 1,
                metadata={"page": index + 1},
            )
        )
    return ParsedDocument(
        pages=pages,
        metadata={"page_count": len(pages), "type": "pdf"},
    )


def parse_docx(file_path: Path) -> ParsedDocument:
    doc = DocxDocument(str(file_path))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    text = "\n".join(paragraphs)
    page = ParsedPage(text=text, page_number=None, metadata={"page": None})
    return ParsedDocument(
        pages=[page],
        metadata={"paragraphs": len(paragraphs), "type": "docx"},
    )


def parse_txt(file_path: Path) -> ParsedDocument:
    text = file_path.read_text(encoding="utf-8", errors="ignore")
    page = ParsedPage(text=text, page_number=None, metadata={"page": None})
    return ParsedDocument(pages=[page], metadata={"type": "txt"})


def parse_md(file_path: Path) -> ParsedDocument:
    text = file_path.read_text(encoding="utf-8", errors="ignore")
    headings = re.findall(r"^#{1,6}\s+(.+)$", text, re.MULTILINE)
    page = ParsedPage(text=text, page_number=None, metadata={"page": None})
    return ParsedDocument(
        pages=[page],
        metadata={"type": "md", "headings": headings[:20]},
    )


def parse_document(file_path: Path, enable_ocr: bool = False) -> ParsedDocument:
    ext = file_path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}")
    if ext == ".pdf":
        return parse_pdf(file_path, enable_ocr)
    if ext == ".docx":
        return parse_docx(file_path)
    if ext == ".txt":
        return parse_txt(file_path)
    if ext == ".md":
        return parse_md(file_path)
    raise ValueError(f"Unsupported file type: {ext}")
